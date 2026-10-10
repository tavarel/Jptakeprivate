"""JPX capital-plan lists."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from radar.ingest.jpx import list_month_from_sheet, load_capital_plan_lists, published_on

HEADER = [
    "業種\nコード",
    "業種",
    "市場区分",
    "証券\nコード",
    "銘柄名",
    "開示状況",
    "開示内容の\nアップデート日",
]
SUBHEADER = [None, None, None, None, None, "要請に基づく\n開示状況", None]


def _sheet(book: Workbook, title: str, rows: list[list[object]], first: bool = False) -> None:
    sheet = book.active if first else book.create_sheet(title)
    assert sheet is not None
    sheet.title = title
    sheet.append([])
    sheet.append(["「資本コストや株価を意識した経営の実現に向けた対応」に関する開示企業一覧表"])
    sheet.append([])
    sheet.append(["2025-06-30", "時点"])
    sheet.append([])
    sheet.append(HEADER)
    sheet.append(SUBHEADER)
    for row in rows:
        sheet.append(row)


def _write_workbook(path: Path) -> None:
    book = Workbook()
    _sheet(
        book,
        "開示企業一覧（2025年6月末時点）  ",  # full-width brackets + trailing spaces, as published
        [
            [3050, "食料品", "スタンダード", 9999, "サンプル工業", "開示済", None],
            [6050, "卸売業", "グロース", "130A", "サンプル商事", "検討中", None],
        ],
        first=True,
    )
    _sheet(
        book,
        "【過去分】開示企業一覧（2025年5月末時点）",
        [[3050, "食料品", "スタンダード", 9999, "サンプル工業", "検討中", None]],
    )
    book.create_sheet("本資料の目的")  # unrelated sheet: ignored
    book.save(path)


@pytest.mark.parametrize(
    ("sheet", "expected"),
    [
        ("開示企業一覧 (2025年6月末時点)", date(2025, 6, 1)),
        ("【過去分】開示企業一覧 (2023年12月末時点)", date(2023, 12, 1)),
        ("【過去分】開示企業一覧（2024年3月末時点）", date(2024, 3, 1)),  # full-width brackets
        ("注記", None),
    ],
)
def test_list_month_from_sheet(sheet: str, expected: date | None) -> None:
    assert list_month_from_sheet(sheet) == expected


def test_published_on_is_conservative() -> None:
    # The list for June is published around 15 July; we only trust it from 1 August.
    assert published_on(date(2025, 6, 1)) == date(2025, 8, 1)
    assert published_on(date(2025, 12, 1)) == date(2026, 2, 1)


def test_load_capital_plan_lists(tmp_path: Path) -> None:
    _write_workbook(tmp_path / "2025-06.xlsx")
    lists = load_capital_plan_lists(tmp_path).sort_values(["list_month", "sec_code"])
    assert lists.columns.tolist() == [
        "list_month",
        "published_on",
        "sec_code",
        "company_name",
        "market_segment",
        "status",
    ]
    assert lists["list_month"].tolist() == [date(2025, 5, 1), date(2025, 6, 1), date(2025, 6, 1)]
    # 4-character JPX codes become EDINET / J-Quants 5-character codes.
    assert lists["sec_code"].tolist() == ["99990", "130A0", "99990"]
    assert lists["status"].tolist() == ["検討中", "検討中", "開示済"]


def test_load_capital_plan_lists_dedupes_overlapping_files(tmp_path: Path) -> None:
    _write_workbook(tmp_path / "2025-06.xlsx")
    _write_workbook(tmp_path / "2025-07.xlsx")  # a newer download repeats the past sheets
    assert len(load_capital_plan_lists(tmp_path)) == 3


def test_load_capital_plan_lists_empty_dir(tmp_path: Path) -> None:
    assert load_capital_plan_lists(tmp_path).empty
